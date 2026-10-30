"""Traceable exports; template writes require an explicit, unique cell map."""
from datetime import datetime
from pathlib import Path
import math
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .core import ASSAYS, Result, file_sha256, normalize_position


HEADERS = ['Rapport', 'Posisjon', 'Prøveinfo', 'Analyse', 'Forventet bp',
           'Peak bp', 'RFU', 'Fortynning (forslag)', 'Status', 'Merknad',
           'Råhøyde', 'RFU-faktor', 'Tabellside', 'Overskriftside',
           'Søkevindu ±bp', 'Manuelt analysevalg', 'Kildefil', 'Metode',
           'Kontrollert av', 'Kontrollkommentar', 'Valgt mål-bp',
           'Kandidater i vindu', 'Kandidat bp / råhøyde', 'Kilde SHA256', 'Avlest PDF-side']


def text_cell(cell, value):
    if isinstance(value, float) and not math.isfinite(value):
        value = str(value)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'  # PDF text is data, including strings beginning '='.


def append(ws, values):
    row = ws.max_row + 1 if ws['A1'].value is not None else 1
    for col, value in enumerate(values, 1):
        text_cell(ws.cell(row, col), value)


def style(ws):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='124E66')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 30
    for col in ws.columns:
        width = min(65, max(13, max(len(str(c.value or '')) for c in col) + 2))
        ws.column_dimensions[get_column_letter(col[0].column)].width = width


def save_new(wb, output: Path):
    output = Path(output).resolve()
    if output.suffix.lower() != '.xlsx':
        raise ValueError('Resultatet må lagres som .xlsx.')
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation also prevents races and accidental overwriting.
    try:
        with output.open('xb') as stream:
            try:
                wb.save(stream)
            except BaseException:
                stream.close()
                output.unlink(missing_ok=True)
                raise
    finally:
        wb.close()


def export_workbook(results: list[Result], output: Path, *, errors=(), approved_only=False):
    if not results and not errors:
        raise ValueError('Ingen resultater å eksportere.')
    total_count = len(results)
    if approved_only:
        results = [r for r in results if r.status in ('KLAR', 'KONTROLLERT') and r.rfu is not None]
        if not results and not errors:
            raise ValueError('Ingen avklarte resultater. Eksporter alle for å se kontrollpunktene.')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resultater'
    append(ws, HEADERS)
    for result in results:
        s, p = result.sample, result.peak
        append(ws, [Path(s.source).name, s.position, s.info, result.assay,
                    result.expected, p.bp if p else None, result.rfu,
                    result.dilution, result.status, result.note,
                    p.height if p else None, result.factor,
                    s.table_page or (p.page if p else None), s.header_page, result.tolerance,
                    'Ja' if result.override else 'Nei', s.source,
                    p.method if p else None, result.reviewer, result.review_note,
                    result.target_bp, len(result.candidates),
                    '; '.join(f'{c.bp:g} bp / {c.height if c.height is not None else "mangler høyde"}'
                              for c in result.candidates), s.source_sha256, p.page if p else None])
        for col in (5, 6, 7, 21):
            ws.cell(ws.max_row, col).number_format = '#,##0.##'
        ws.cell(ws.max_row, 11).number_format = '0.000'
        if result.status == 'KONTROLL':
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill('solid', fgColor='FFF0D5')
    style(ws)
    peaks = wb.create_sheet('Alle topper')
    append(peaks, ['Rapport', 'Posisjon', 'Prøveinfo', 'Peak bp', 'Råhøyde', 'Side', 'Kildefil', 'Metode'])
    for result in results:
        for p in result.sample.peaks:
            append(peaks, [Path(result.sample.source).name, result.sample.position,
                           result.sample.info, p.bp, p.height, p.page,
                           result.sample.source, p.method])
    style(peaks)
    settings = wb.create_sheet('Innstillinger')
    append(settings, ['Innstilling', 'Verdi'])
    for row in [('Programversjon', __version__), ('Eksportert lokalt', datetime.now().isoformat(timespec='seconds')),
                ('RFU-faktorer', ', '.join(str(x) for x in sorted({r.factor for r in results}, key=str))),
                ('Eksportutvalg', 'Bare avklarte' if approved_only else 'Alle resultater'),
                ('Utelatte prøver', total_count - len(results)),
                ('RFU-valg', 'Eksplisitt faktor valgt av bruker. None betyr ubekreftet. Tusentall alene bekrefter ikke faktor.'),
                ('Metode', 'Height fra PDF-teksttabell; manuell grafavlesning merkes separat.'),
                ('Valg av topp', 'Én kandidat: automatisk. Flere: ingen RFU før eksakt bp velges; krever kontroll.'),
                ('Grense 1500', '1500–<2000: 1:2'), ('Grense 2000', '2000–<4000: 1:5'),
                ('Grense 4000', '4000–6000: 1:10'), ('Under 1500', 'Ufortynnet'),
                ('Over 6000', '1:15 eller 1:20; krever laboratorievalg'),
                ('Grensenær verdi', 'Innen ±10 % av 1500, 2000, 4000 eller 6000 krever kontroll'),
                ('Bruk', 'Beslutningsstøtte. Ikke klinisk validert. Kontroller skala og regler før rutinebruk.')]:
        append(settings, row)
    for name, bp in ASSAYS.items():
        append(settings, (name, bp))
    style(settings)
    if errors:
        failures = wb.create_sheet('Filfeil')
        append(failures, ['Fil', 'Feil'])
        for row in errors:
            append(failures, row)
        style(failures)
    save_new(wb, output)


def fill_template(results: list[Result], template: Path, profile: dict, output: Path):
    template, output = Path(template).resolve(), Path(output).resolve()
    if template == output or any(Path(r.sample.source).resolve() == output for r in results):
        raise ValueError('Originalfilen kan ikke overskrives.')
    if template.suffix.lower() != '.xlsx':
        raise ValueError('Kun vanlige .xlsx-maler støttes i første versjon.')
    by_key = {}
    for r in results:
        key = (Path(r.sample.source).name.casefold(), r.sample.position)
        if key in by_key:
            raise ValueError(f'Flere resultater for {key}; velg én rapportkjøring om gangen.')
        by_key[key] = r
        if r.sample.source_sha256 and (not Path(r.sample.source).is_file() or
                                      file_sha256(r.sample.source) != r.sample.source_sha256):
            raise ValueError('Kildefilen er flyttet eller endret. Les rapporten på nytt før maleksport.')
    if not isinstance(profile.get('rows'), list) or not profile['rows']:
        raise ValueError('Malprofilen må inneholde en ikke-tom rows-liste.')
    wb = load_workbook(template)
    try:
        if profile.get('sheet') not in wb.sheetnames:
            raise ValueError('Arknavnet i malprofilen finnes ikke i malen.')
        ws = wb[profile['sheet']]
        used = set()
        used_keys = set()
        for mapping in profile['rows']:
            key = (mapping['report'].casefold(), normalize_position(mapping['position']))
            if key not in by_key or key in used_keys:
                raise ValueError(f'Manglende eller gjentatt prøve i malprofil: {key}')
            used_keys.add(key)
            r = by_key[key]
            if r.status not in ('KLAR', 'KONTROLLERT') or r.rfu is None or not r.dilution:
                raise ValueError(f'{key}: resultatet krever kontroll før innfylling i mal.')
            values = {'rfu': r.rfu, 'bp': r.peak.bp, 'dilution': r.dilution,
                      'assay': r.assay, 'position': r.sample.position}
            cells = mapping.get('cells', {})
            if not cells:
                raise ValueError('Hver prøve må ha minst én målcelle.')
            for field, address in cells.items():
                if field not in values or not isinstance(address, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', address):
                    raise ValueError(f'Ugyldig felt/celle: {field}: {address}')
                if any(address in merged for merged in ws.merged_cells.ranges):
                    raise ValueError(f'Celle {address} er sammenslått; velg en vanlig målcelle.')
                if ws[address].column > 16384:
                    raise ValueError('Målcelle ligger utenfor Excels kolonnegrense.')
                if address in used or ws[address].value is not None:
                    raise ValueError(f'Celle {address} er opptatt eller brukes flere ganger.')
                used.add(address)
                text_cell(ws[address], values[field])
        # All checks succeed before a new file is created. The original is untouched.
        save_new(wb, output)
    finally:
        wb.close()
