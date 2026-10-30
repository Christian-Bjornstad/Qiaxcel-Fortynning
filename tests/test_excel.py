from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from test_core import Sample, Peak, analyze
from qiaxcel.excel import export_workbook, fill_template


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.result = analyze(Sample('report.pdf', 'A01', 'TERT', 1,
                                     [Peak(192, 3.1, 2)]), factor=1000)

    def test_export_roundtrip_and_formula_text(self):
        self.result.sample.info = '=HYPERLINK("bad")'
        out = self.root / 'out.xlsx'
        export_workbook([self.result], out)
        wb = load_workbook(out)
        self.addCleanup(wb.close)
        self.assertEqual(wb['Resultater']['G2'].value, 3100)
        self.assertEqual(wb['Resultater']['H2'].value, '1:5')
        self.assertEqual(wb['Resultater']['C2'].data_type, 's')
        self.assertIn('Alle topper', wb.sheetnames)
        self.assertIn('Innstillinger', wb.sheetnames)
        with self.assertRaises(FileExistsError):
            export_workbook([self.result], out)

    def profile(self, destination='C2'):
        template = self.root / 'template.xlsx'
        wb = Workbook()
        wb.active.title = 'Oppsett'
        wb.active['A2'] = 'A01'
        wb.active['D2'] = '=1+1'
        wb.save(template)
        wb.close()
        profile = {'sheet': 'Oppsett', 'rows': [
            {'report': 'report.pdf', 'position': 'A1',
             'cells': {'rfu': 'B2', 'dilution': destination}}]}
        return template, profile

    def test_template_copy_preserves_original_and_unrelated_formula(self):
        template, profile = self.profile()
        before = template.read_bytes()
        out = self.root / 'filled.xlsx'
        fill_template([self.result], template, profile, out)
        self.assertEqual(template.read_bytes(), before)
        wb = load_workbook(out)
        self.addCleanup(wb.close)
        self.assertEqual(wb['Oppsett']['B2'].value, 3100)
        self.assertEqual(wb['Oppsett']['D2'].value, '=1+1')

    def test_reject_occupied_cell_duplicate_result_and_unreviewed(self):
        template, profile = self.profile(destination='D2')
        out = self.root / 'filled.xlsx'
        with self.assertRaises(ValueError):
            fill_template([self.result], template, profile, out)
        template, profile = self.profile()
        with self.assertRaises(ValueError):
            fill_template([self.result, self.result], template, profile, out)
        self.result.status = 'KONTROLL'
        with self.assertRaises(ValueError):
            fill_template([self.result], template, profile, out)
        self.assertFalse(out.exists())

    def test_input_cannot_be_overwritten(self):
        template, profile = self.profile()
        with self.assertRaises(ValueError):
            fill_template([self.result], template, profile, template)

    def test_merged_cells_are_rejected(self):
        template, profile = self.profile()
        wb = load_workbook(template)
        wb['Oppsett'].merge_cells('B2:C2')
        wb.save(template)
        wb.close()
        with self.assertRaises(ValueError):
            fill_template([self.result], template, profile, self.root / 'out.xlsx')

    def test_ambiguous_rfu_is_blank_and_selected_bp_is_traceable(self):
        sample = Sample('demo.pdf', 'A01', 'TERT', 1,
                        [Peak(190, 3.1, 2), Peak(201, 8.9, 2)], source_sha256='synthetic-hash')
        ambiguous = analyze(sample, factor=1000)
        out = self.root / 'ambiguous.xlsx'
        export_workbook([ambiguous], out)
        wb = load_workbook(out)
        self.addCleanup(wb.close)
        self.assertIsNone(wb['Resultater']['G2'].value)
        headings = {c.value: c.column for c in wb['Resultater'][1]}
        self.assertEqual(wb['Resultater'].cell(2, headings['Kandidater i vindu']).value, 2)
        self.assertEqual(wb['Resultater'].cell(2, headings['Kilde SHA256']).value, 'synthetic-hash')
        selected = analyze(sample, factor=1000, target_bp=190)
        selected_out = self.root / 'selected.xlsx'
        export_workbook([selected], selected_out)
        wb2 = load_workbook(selected_out)
        self.addCleanup(wb2.close)
        self.assertEqual(wb2['Resultater'].cell(2, headings['Valgt mål-bp']).value, 190)
        self.assertEqual(wb2['Resultater']['G2'].value, 3100)

    def test_approved_export_excludes_pending_but_keeps_file_errors(self):
        pending = analyze(Sample('pending.pdf', 'A02', 'TERT', 1,
                                 [Peak(190, 1, 2), Peak(200, 4, 2)]), factor=1000)
        out = self.root / 'approved.xlsx'
        export_workbook([self.result, pending], out, approved_only=True,
                        errors=[('bad.pdf', 'Synthetic file error')])
        wb = load_workbook(out)
        self.addCleanup(wb.close)
        self.assertEqual(wb['Resultater'].max_row, 2)
        self.assertEqual(wb['Filfeil']['A2'].value, 'bad.pdf')

    def test_template_rejects_stale_file_hash(self):
        template, profile = self.profile()
        pdf = self.root / 'report.pdf'
        pdf.write_bytes(b'changed synthetic data')
        self.result.sample.source = str(pdf)
        self.result.sample.source_sha256 = 'wrong-old-hash'
        with self.assertRaises(ValueError):
            fill_template([self.result], template, profile, self.root / 'filled.xlsx')


if __name__ == '__main__':
    unittest.main()
