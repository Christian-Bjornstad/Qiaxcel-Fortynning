import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from PyQt6.QtWidgets import QApplication, QDialog, QMessageBox
from PyQt6.QtCore import Qt, QEventLoop, QTimer
from openpyxl import load_workbook
from qiaxcel.core import Sample, Peak, analyze, review_result
from qiaxcel.gui import App
from PyQt6.QtWidgets import QMainWindow

class MigrationTests(unittest.TestCase):
    def test_application_is_native_qt_window(self):
        self.assertTrue(issubclass(App, QMainWindow))


class GuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.qt = QApplication.instance() or QApplication([])

    def setUp(self):
        self.app = App()
        self.addCleanup(self.cleanup_window)

    def cleanup_window(self):
        self.app.unsaved_changes = False
        self.app.close()
        self.app.deleteLater()
        self.qt.processEvents()

    def sample(self, peaks=None, info='TERT'):
        return Sample('synthetic.pdf', 'A01', info, 1, peaks or [Peak(192, 3.1, 2)])

    def select(self, row=0):
        self.app.table.selectRow(row)
        self.qt.processEvents()

    def test_ambiguous_candidates_do_not_show_rfu_until_exact_selection(self):
        self.app.rows = [analyze(self.sample([Peak(180, 2.5, 2), Peak(192, 4.5, 3)]), factor=1000)]
        self.app.refresh()
        self.select()
        self.assertEqual(self.app.table.item(0, 5).text(), '—')
        self.assertEqual(self.app.candidate_table.rowCount(), 2)
        self.app.candidate_table.selectRow(0)
        self.app.pick_candidate()
        self.assertEqual(self.app.rows[0].peak.bp, 180)
        self.assertEqual(self.app.rows[0].rfu, 2500)
        self.assertEqual(self.app.table.item(0, 5).text(), '2 500')

    def test_missing_height_stays_blank_and_candidate_is_preserved(self):
        self.app.rows = [analyze(self.sample([Peak(192, None, 2)]), factor=1000)]
        self.app.refresh()
        self.select()
        self.assertEqual(self.app.table.item(0, 4).text(), '192')
        self.assertEqual(self.app.table.item(0, 5).text(), '—')
        self.assertEqual(self.app.candidate_table.item(0, 1).text(), 'Mangler høyde')

    def test_cancel_settings_keeps_signed_review_and_applied_scale(self):
        self.app.active_settings = {'factor': 1000, 'tolerance': 20}
        original = analyze(self.sample(), factor=1000)
        signed = review_result(original, assay='TERT', bp=192, raw_height=3.1,
                               reviewer='TEST', comment='Synthetic review', chosen_dilution='1:5')
        self.app.rows = [signed]
        self.app.refresh()
        self.app.scale_box.setCurrentIndex(0)
        with patch.object(QMessageBox, 'warning', return_value=QMessageBox.StandardButton.Cancel):
            self.app.apply_settings()
        self.assertIs(self.app.rows[0], signed)
        self.assertEqual(self.app.active_settings['factor'], 1000)
        self.assertEqual(self.app.scale_box.currentData(), 1000)

    def test_cancel_reimport_does_not_remove_reviews_or_errors(self):
        signed = review_result(analyze(self.sample(), factor=1000), assay='TERT', bp=192,
            raw_height=3.1, reviewer='TEST', comment='Synthetic review', chosen_dilution='1:5')
        self.app.rows = [signed]
        self.app.errors = [('bad.pdf', 'Unreadable synthetic report')]
        self.app.paths = [Path('new.pdf')]
        with patch.object(QMessageBox, 'warning', return_value=QMessageBox.StandardButton.Cancel):
            self.app.read()
        self.assertEqual(self.app.rows, [signed])
        self.assertEqual(len(self.app.errors), 1)
        self.assertFalse(self.app.busy)

    def test_search_and_bulk_analysis_apply_only_selected_rows(self):
        self.app.rows = [analyze(self.sample(info='first'), factor=1000),
                         analyze(Sample('synthetic.pdf', 'B02', 'second', 1, [Peak(270, 3, 2)]), factor=1000)]
        self.app.refresh()
        self.select(0)
        self.app.assay_box.setCurrentText('TERT')
        self.app.set_assay()
        self.assertEqual(self.app.rows[0].assay, 'TERT')
        self.assertIsNone(self.app.rows[1].assay)
        self.app.search_edit.setText('second')
        self.assertTrue(self.app.table.isRowHidden(0))
        self.assertFalse(self.app.table.isRowHidden(1))

    def test_review_dialog_saves_explicit_signoff(self):
        result = analyze(self.sample(), factor=1000)
        from qiaxcel.gui import ReviewDialog
        dialog = ReviewDialog(result, self.app)
        self.addCleanup(dialog.deleteLater)
        dialog.reviewer_edit.setText('TEST')
        dialog.comment_edit.setPlainText('Synthetic review')
        dialog.dilution_box.setCurrentText('1:5')
        dialog.save_review()
        self.assertEqual(dialog.result_value.status, 'KONTROLLERT')
        self.assertEqual(dialog.result_value.rfu, 3100)
        self.assertEqual(dialog.result_value.reviewer, 'TEST')

    def test_failed_pdf_remains_visible_and_exportable(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / 'broken.pdf'
            bad.write_text('not a PDF', encoding='utf8')
            self.app.paths = [bad]
            self.app.read()
            loop = QEventLoop()
            timer = QTimer()
            timer.setInterval(10)
            timer.timeout.connect(lambda: loop.quit() if not self.app.busy else None)
            timer.start()
            QTimer.singleShot(5000, loop.quit)
            loop.exec()
            timer.stop()
            self.assertFalse(self.app.busy)
            self.assertEqual(len(self.app.errors), 1)
            self.assertIn('broken.pdf', self.app.error_list.item(0).text())
            output = Path(tmp) / 'result.xlsx'
            with patch.object(self.app, 'output_path', return_value=str(output)):
                self.assertTrue(self.app.export())
            wb = load_workbook(output)
            self.addCleanup(wb.close)
            self.assertIn('Filfeil', wb.sheetnames)
            self.assertIn('broken.pdf', str(list(wb['Filfeil'].values)))

    def test_approved_export_preserves_failures_and_excludes_unresolved_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.app.rows = [analyze(self.sample([Peak(192, 3, 2)]), factor=1000),
                             analyze(self.sample([Peak(180, 2, 2), Peak(192, 4, 2)]), factor=1000)]
            self.app.errors = [('bad.pdf', 'Synthetic failure')]
            self.app.export_scope.setCurrentIndex(1)
            output = Path(tmp) / 'approved.xlsx'
            with patch.object(self.app, 'output_path', return_value=str(output)):
                self.assertTrue(self.app.export())
            wb = load_workbook(output)
            self.addCleanup(wb.close)
            self.assertEqual(wb['Resultater'].max_row, 2)
            self.assertIn('Synthetic failure', str(list(wb['Filfeil'].values)))


    def test_graph_calibration_converts_actual_rfu_back_to_raw_once(self):
        from qiaxcel.qt_preview import GraphDialog, PdfPreview
        from PyQt6.QtGui import QImage
        from PyQt6.QtCore import QBuffer, QByteArray, QIODevice
        result = analyze(self.sample([Peak(192, None, 2)]), factor=1000)
        with patch.object(PdfPreview, 'load_page'):
            dialog = GraphDialog(result, self.app)
        self.addCleanup(dialog.deleteLater)
        image = QImage(600, 400, QImage.Format.Format_RGB32)
        image.fill(Qt.GlobalColor.white)
        data = QByteArray()
        buffer = QBuffer(data)
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        image.save(buffer, 'PNG')
        dialog.preview.png = bytes(data)
        dialog.preview.loaded_identity = dialog.preview.identity
        dialog.preview.view.set_image(bytes(data))
        dialog.pick_point(100, 300)
        dialog.pick_point(100, 100)
        dialog.pick_point(200, 200)
        self.assertEqual(dialog.raw_height, 2.5)
        self.assertTrue(dialog.use_button.isEnabled())
        self.assertIn('2500', dialog.calibration_note)
        self.assertIn('300', dialog.calibration_note)

    def test_close_cancellation_keeps_unsaved_signed_review(self):
        self.app.rows = [review_result(analyze(self.sample(), factor=1000), assay='TERT',
            bp=192, raw_height=3.1, reviewer='TEST', comment='Synthetic', chosen_dilution='1:5')]
        self.app.unsaved_changes = True
        self.app.show()
        with patch.object(QMessageBox, 'warning', return_value=QMessageBox.StandardButton.Cancel):
            self.app.close()
        self.assertTrue(self.app.isVisible())
        self.assertEqual(self.app.rows[0].reviewer, 'TEST')

    def test_graph_review_keeps_calibration_context_beside_human_comment(self):
        from qiaxcel.gui import ReviewDialog
        result = analyze(self.sample([Peak(192, None, 2)]), factor=1000)
        dialog = ReviewDialog(result, self.app, raw_height=2.5, source_page=2,
                              method='Kalibrert manuell grafavlesning',
                              calibration_note='Syntetisk kalibrering 2500 RFU')
        self.addCleanup(dialog.deleteLater)
        dialog.reviewer_edit.setText('TEST')
        dialog.comment_edit.setPlainText('Kontrollert toppetikett')
        dialog.dilution_box.setCurrentText('1:5')
        dialog.save_review()
        self.assertEqual(dialog.result_value.rfu, 2500)
        self.assertIn('Kontrollert toppetikett', dialog.result_value.review_note)
        self.assertIn('2500 RFU', dialog.result_value.review_note)

    def test_pdf_preview_rejects_changed_source_and_renders_original(self):
        from qiaxcel.qt_preview import PageRenderer
        from qiaxcel.core import file_sha256
        from PyQt6.QtTest import QSignalSpy
        from PyQt6.QtGui import QImage
        from test_integration import make_pdf
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / 'synthetic.pdf'
            make_pdf(pdf, ['Synthetic source page'])
            original_hash = file_sha256(pdf)
            worker = PageRenderer(pdf, 1, original_hash)
            rendered = QSignalSpy(worker.rendered)
            worker.run()
            self.assertEqual(len(rendered), 1)
            self.assertFalse(QImage.fromData(rendered[0][0]).isNull())
            pdf.write_bytes(pdf.read_bytes() + b'\n%changed\n')
            worker = PageRenderer(pdf, 1, original_hash)
            failures = QSignalSpy(worker.failed)
            rendered = QSignalSpy(worker.rendered)
            worker.run()
            self.assertEqual(len(rendered), 0)
            self.assertIn('endret', failures[0][0])

    def test_graph_reload_invalidates_previous_measurement_before_render(self):
        from qiaxcel.qt_preview import GraphDialog, PageRenderer, PdfPreview
        result = analyze(self.sample([Peak(192, None, 2)]), factor=1000)
        with patch.object(PdfPreview, 'load_page'):
            dialog = GraphDialog(result, self.app)
        self.addCleanup(dialog.deleteLater)
        dialog.raw_height = 2.5
        dialog.use_button.setEnabled(True)
        dialog.preview.png = b'old image'
        # Slow external rendering is replaced, while the real loading lifecycle runs.
        with patch.object(PageRenderer, 'run', lambda worker: worker.finished.emit()):
            dialog.preview.load_page()
            invalidated = not dialog.use_button.isEnabled() and dialog.raw_height is None
            loop = QEventLoop()
            timer = QTimer()
            timer.setInterval(10)
            timer.timeout.connect(lambda: loop.quit() if not dialog.preview.is_busy() else None)
            timer.start()
            QTimer.singleShot(2000, loop.quit)
            loop.exec()
            timer.stop()
        self.assertTrue(invalidated)

    def test_review_roundtrips_raw_value_below_dilution_threshold(self):
        from qiaxcel.gui import ReviewDialog
        result = analyze(self.sample([Peak(192, 3.999999, 2)]), factor=1000)
        dialog = ReviewDialog(result, self.app)
        self.addCleanup(dialog.deleteLater)
        self.assertEqual(float(dialog.raw_edit.text()), 3.999999)
        dialog.reviewer_edit.setText('TEST')
        dialog.comment_edit.setPlainText('Synthetic threshold review')
        dialog.save_review()
        self.assertAlmostEqual(dialog.result_value.rfu, 3999.999)
        self.assertEqual(dialog.result_value.dilution, '1:5')

    def test_review_roundtrips_exact_fractional_candidate_bp(self):
        from qiaxcel.gui import ReviewDialog
        result = analyze(self.sample([Peak(192.123456789, 3.1, 2)]), factor=1000)
        dialog = ReviewDialog(result, self.app)
        self.addCleanup(dialog.deleteLater)
        self.assertEqual(float(dialog.bp_edit.text()), 192.123456789)
        dialog.reviewer_edit.setText('TEST')
        dialog.comment_edit.setPlainText('Synthetic exact bp review')
        dialog.save_review()
        self.assertEqual(dialog.result_value.peak.bp, 192.123456789)

    def test_measured_field_edit_requires_new_dilution_choice(self):
        from qiaxcel.gui import ReviewDialog
        dialog = ReviewDialog(analyze(self.sample(), factor=1000), self.app)
        self.addCleanup(dialog.deleteLater)
        self.assertEqual(dialog.dilution_box.currentText(), '1:5')
        dialog.raw_edit.setText('4')
        self.assertEqual(dialog.dilution_box.currentIndex(), 0)

    def test_graph_applied_value_requires_new_dilution_choice(self):
        from qiaxcel.gui import ReviewDialog
        dialog = ReviewDialog(analyze(self.sample(), factor=1000), self.app,
                              raw_height=5, method='Kalibrert manuell grafavlesning')
        self.addCleanup(dialog.deleteLater)
        self.assertEqual(dialog.dilution_box.currentIndex(), 0)

    def test_pdf_page_change_immediately_removes_old_image(self):
        from qiaxcel.qt_preview import PdfPreview
        from PyQt6.QtGui import QImage
        from PyQt6.QtCore import QBuffer, QByteArray, QIODevice
        preview = PdfPreview(self.app)
        preview.set_source('synthetic.pdf', 1, 'synthetic-hash')
        image = QImage(600, 400, QImage.Format.Format_RGB32)
        image.fill(Qt.GlobalColor.white)
        data = QByteArray()
        buffer = QBuffer(data)
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        image.save(buffer, 'PNG')
        preview.png = bytes(data)
        preview.view.set_image(bytes(data))
        preview.page_box.setValue(2)
        self.assertIsNone(preview.png)
        self.assertIsNone(preview.view.image_item)

    def test_pdf_default_zoom_uses_width_so_table_is_readable(self):
        from qiaxcel.qt_preview import PageView
        from PyQt6.QtGui import QImage
        from PyQt6.QtCore import QBuffer, QByteArray, QIODevice
        view = PageView()
        self.addCleanup(view.close)
        self.addCleanup(view.deleteLater)
        view.resize(450, 260)
        view.show()
        self.qt.processEvents()
        image = QImage(900, 1300, QImage.Format.Format_RGB32)
        image.fill(Qt.GlobalColor.white)
        data = QByteArray()
        buffer = QBuffer(data)
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        image.save(buffer, 'PNG')
        view.set_image(bytes(data))
        displayed_width = view.transform().m11() * 900
        self.assertGreater(displayed_width, view.viewport().width() * 0.9)

    def test_pdf_new_source_hash_ignores_stale_render_completion(self):
        from qiaxcel.qt_preview import PdfPreview
        preview = PdfPreview(self.app)
        preview.set_source('synthetic.pdf', 1, 'first-hash')
        preview._render_identity = preview.identity
        preview.png = b'old picture'
        preview.loaded_identity = preview.identity
        preview.set_source('synthetic.pdf', 1, 'second-hash')
        preview._show_page(b'stale picture', 2)
        self.assertIsNone(preview.png)
        self.assertFalse(preview.has_current_image())

    def test_graph_cannot_accept_a_measurement_for_a_different_page(self):
        from qiaxcel.qt_preview import GraphDialog, PdfPreview
        with patch.object(PdfPreview, 'load_page'):
            dialog = GraphDialog(analyze(self.sample(), factor=1000), self.app)
        self.addCleanup(dialog.deleteLater)
        dialog.preview.png = b'stale picture'
        dialog.preview.loaded_identity = ('synthetic.pdf', 1, '')
        dialog.preview.page_box.setValue(2)
        dialog.raw_height = 2.5
        dialog.accept()
        self.assertNotEqual(dialog.result(), QDialog.DialogCode.Accepted)


if __name__ == '__main__':
    unittest.main()
